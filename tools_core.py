"""
tools_core.py — 7 Core Tools with Safety Layer
Unified tool system for the UniVoid AI assistant.

Tools:
  1. terminal_executor   — Run shell commands (with safety blocklist)
  2. file_manager        — Create/read/write/delete/move/list/search files
  3. web_search          — DuckDuckGo search (no API key needed)
  4. document_creator    — Create txt/md/html/py/docx files
  5. app_opener          — Launch applications by name
  6. system_info         — OS, disk, processes, environment
  7. memory_store        — Persistent user memory
"""

import os
import sys
import json
import shutil
import platform
import subprocess
import psutil
import requests
import memory as mem
from config import DESKTOP, HOME, WORKSPACE

# Ensure WORKSPACE exists
os.makedirs(WORKSPACE, exist_ok=True)


# ============================================
# SAFETY LAYER
# ============================================

BLOCKED_COMMANDS = [
    # Destructive disk operations
    "rm -rf /", "rm -rf /*", "rm -rf ~",
    "del /f /s /q c:\\", "del /f /s /q c:/",
    "format c:", "format c:\\",
    "rd /s /q c:\\", "rmdir /s /q c:\\",
    # System destruction
    "shutdown /s", "shutdown -s",
    "shutdown /r", "shutdown -r",
    ":(){:|:&};:",  # Fork bomb
    "mkfs", "dd if=/dev/zero",
    # Registry destruction
    "reg delete hklm", "reg delete hkcu",
    # Network attacks
    "netsh advfirewall set allprofiles state off",
]

BLOCKED_PATTERNS = [
    "format c", "format d", "format e",
    "del /f /s /q c:", "del /f /s /q d:",
    "rm -rf /", "rm -rf /*",
    "rd /s /q c:", "rd /s /q d:",
    "reg delete hk",
    "cipher /w:c:",
    "diskpart",
    "bcdedit",
    "bootrec",
]


def _is_command_safe(command):
    """Check if a command is safe to execute."""
    cmd_lower = command.lower().strip()

    # Check exact blocked commands
    for blocked in BLOCKED_COMMANDS:
        if blocked in cmd_lower:
            return False, f"🛑 BLOCKED: '{command}' is a destructive command."

    # Check blocked patterns
    for pattern in BLOCKED_PATTERNS:
        if pattern in cmd_lower:
            return False, f"🛑 BLOCKED: Command matches dangerous pattern '{pattern}'."

    return True, "OK"


# ============================================
# TOOL 1: TERMINAL EXECUTOR
# ============================================

def terminal_executor(args):
    """
    Run a shell command and return the output.
    Args: {"command": "pip install flask"}
    """
    command = args.get("command", "").strip()
    if not command:
        return "❌ No command provided."

    # Safety check
    safe, reason = _is_command_safe(command)
    if not safe:
        return reason

    timeout = args.get("timeout", 30)

    try:
        result = subprocess.run(
            command, shell=True,
            capture_output=True, text=True,
            timeout=timeout, cwd=args.get("cwd", None)
        )
        output = result.stdout.strip()
        error = result.stderr.strip()

        response = ""
        if output:
            response += f"✅ Output:\n{output}"
        if error:
            if response:
                response += "\n"
            response += f"⚠️ Stderr:\n{error}"
        if not response:
            response = f"✅ Command completed (exit code {result.returncode})"

        # Truncate long output
        if len(output) > 2000:
            output = output[:2000] + "\n...(truncated)"
            response = f"✅ Output:\n{output}"
        if error and len(error) > 500:
            error = error[:500] + "\n...(truncated)"
            if response:
                response += f"\n⚠️ Stderr:\n{error}"
            else:
                response = f"⚠️ Stderr:\n{error}"

        return response

    except subprocess.TimeoutExpired:
        return f"❌ Command timed out after {timeout}s."
    except Exception as e:
        return f"❌ Error: {e}"


# ============================================
# TOOL 2: FILE MANAGER
# ============================================

def _resolve_path(path):
    """Resolve relative paths to WORKSPACE."""
    if not path:
        return WORKSPACE
    if not os.path.isabs(path):
        return os.path.join(WORKSPACE, path)
    return path


def file_manager(args):
    """
    Unified file operations.
    Args: {"action": "create|read|write|delete|move|copy|list|search|mkdir",
           "path": "...", "content": "...", "destination": "..."}
    """
    action = args.get("action", "").lower()
    path = _resolve_path(args.get("path", ""))

    if action == "create" or action == "write":
        content = args.get("content", "")
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        return f"✅ File created: {path}"

    elif action == "read":
        if not os.path.isfile(path):
            return f"❌ File not found: {path}"
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        if len(content) > 10000:
            return content[:10000] + f"\n\n...(truncated, {len(content)} total chars)"
        return content

    elif action == "delete":
        if not os.path.exists(path):
            return f"❌ Path not found: {path}"
        if os.path.isdir(path):
            shutil.rmtree(path)
            return f"✅ Folder deleted: {path}"
        else:
            os.remove(path)
            return f"✅ File deleted: {path}"

    elif action == "move" or action == "rename":
        dst = _resolve_path(args.get("destination", ""))
        if not dst:
            return "❌ No destination provided."
        shutil.move(path, dst)
        return f"✅ Moved: {path} → {dst}"

    elif action == "copy":
        dst = _resolve_path(args.get("destination", ""))
        if not dst:
            return "❌ No destination provided."
        if os.path.isdir(path):
            shutil.copytree(path, dst)
        else:
            os.makedirs(os.path.dirname(dst), exist_ok=True)
            shutil.copy2(path, dst)
        return f"✅ Copied: {path} → {dst}"

    elif action == "list":
        if not os.path.exists(path):
            return f"❌ Path not found: {path}"
        items = os.listdir(path)
        if not items:
            return f"📂 {path} is empty."
        lines = [f"📂 {path}:"]
        for i in sorted(items):
            full = os.path.join(path, i)
            icon = "📁" if os.path.isdir(full) else "📄"
            size = ""
            if os.path.isfile(full):
                s = os.path.getsize(full)
                if s > 1048576:
                    size = f" ({s / 1048576:.1f} MB)"
                elif s > 1024:
                    size = f" ({s / 1024:.1f} KB)"
                else:
                    size = f" ({s} B)"
            lines.append(f"  {icon} {i}{size}")
        return "\n".join(lines)

    elif action == "search":
        query = args.get("query", args.get("pattern", ""))
        if not query:
            return "❌ No search query provided."
        results = []
        for root, dirs, files in os.walk(path):
            for f in files:
                if query.lower() in f.lower():
                    results.append(os.path.join(root, f))
                    if len(results) >= 20:
                        break
            if len(results) >= 20:
                break
        if not results:
            return f"❌ No files matching '{query}' in {path}"
        return "🔍 Found:\n" + "\n".join(f"  📄 {r}" for r in results)

    elif action == "mkdir":
        os.makedirs(path, exist_ok=True)
        return f"✅ Folder created: {path}"

    else:
        return f"❌ Unknown action: {action}. Use: create, read, write, delete, move, copy, list, search, mkdir"


# ============================================
# TOOL 3: WEB SEARCH
# ============================================

def web_search(args):
    """
    Search the web using DuckDuckGo.
    Args: {"query": "search terms"}
    """
    query = args.get("query", "").strip()
    if not query:
        return "❌ No search query provided."
    
    try:
        # First try DuckDuckGo Instant Answer API
        params = {"q": query, "format": "json"}
        response = requests.get("https://api.duckduckgo.com/", params=params, timeout=10)
        
        if response.status_code == 200:
            data = response.json()
            
            # Check for instant answer
            if data.get("AbstractText"):
                return f"🔍 {data['AbstractText']}\n\nSource: {data.get('AbstractURL', 'DuckDuckGo')}"
            
            # Check for related topics
            if data.get("RelatedTopics"):
                results = []
                for topic in data["RelatedTopics"][:3]:
                    if isinstance(topic, dict) and topic.get("Text"):
                        results.append(f"• {topic['Text']}")
                if results:
                    return f"🔍 Results for '{query}':\n" + "\n".join(results)
        
        # Fallback to duckduckgo_search library
        try:
            from duckduckgo_search import DDGS
            
            with DDGS() as ddgs:
                results = list(ddgs.text(query, max_results=5))
                
                if not results:
                    return f"❌ No results found for '{query}'"
                
                output = [f"🔍 Search results for '{query}':\n"]
                for i, result in enumerate(results, 1):
                    title = result.get('title', 'No title')
                    body = result.get('body', '')
                    link = result.get('href', '')
                    
                    output.append(f"{i}. {title}")
                    if body:
                        # Truncate body
                        body_short = body[:150] + ("..." if len(body) > 150 else "")
                        output.append(f"   {body_short}")
                    if link:
                        output.append(f"   {link}")
                    output.append("")
                
                return "\n".join(output)
        
        except ImportError:
            return "❌ duckduckgo_search library not installed. Run: pip install duckduckgo-search"
        except Exception as e:
            return f"❌ Search error: {e}"
    
    except Exception as e:
        return f"❌ Search failed: {e}"


# ============================================
# TOOL 4: DOCUMENT CREATOR
# ============================================

def document_creator(args):
    """
    Create various document types (txt, md, html, py, docx).
    Args: {"filename": "report.md", "content": "# Title\nContent", "format": "md"}
    """
    filename = args.get("filename", "").strip()
    content = args.get("content", "").strip()
    doc_format = args.get("format", "").lower().strip()
    
    if not filename:
        return "❌ No filename provided."
    
    if not content:
        return "❌ No content provided."
    
    # Infer format from filename if not provided
    if not doc_format:
        ext = os.path.splitext(filename)[1].lstrip(".")
        doc_format = ext if ext else "txt"
    
    # Resolve path (relative to WORKSPACE)
    if not os.path.isabs(filename):
        filepath = os.path.join(WORKSPACE, filename)
    else:
        filepath = filename
    
    # Ensure filename has correct extension
    if not filename.endswith(f".{doc_format}"):
        filepath = filepath.rsplit(".", 1)[0] + f".{doc_format}"
    
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    
    try:
        if doc_format in ("txt", "md", "py", "js", "java", "cpp", "c", "sh", "bat", "ps1"):
            # Plain text files
            with open(filepath, "w", encoding="utf-8") as f:
                f.write(content)
            return f"✅ {doc_format.upper()} file created: {filepath}"
        
        elif doc_format == "html":
            # Generate styled HTML
            html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Document</title>
    <style>
        body {{
            font-family: Arial, sans-serif;
            max-width: 800px;
            margin: 40px auto;
            padding: 20px;
            line-height: 1.6;
            color: #333;
        }}
        h1, h2, h3 {{
            color: #2c3e50;
        }}
    </style>
</head>
<body>
{content}
</body>
</html>"""
            with open(filepath, "w", encoding="utf-8") as f:
                f.write(html_content)
            return f"✅ HTML file created: {filepath}"
        
        elif doc_format == "docx":
            # Create Word document
            try:
                from docx import Document
                
                doc = Document()
                
                # Split content by lines and add paragraphs
                for line in content.split("\n"):
                    if line.strip().startswith("# "):
                        # Heading 1
                        doc.add_heading(line.strip()[2:], level=1)
                    elif line.strip().startswith("## "):
                        # Heading 2
                        doc.add_heading(line.strip()[3:], level=2)
                    elif line.strip().startswith("### "):
                        # Heading 3
                        doc.add_heading(line.strip()[4:], level=3)
                    else:
                        # Regular paragraph
                        doc.add_paragraph(line)
                
                doc.save(filepath)
                return f"✅ DOCX file created: {filepath}"
            
            except ImportError:
                return "❌ python-docx library not installed. Run: pip install python-docx"
        
        else:
            # Unknown format - save as text
            with open(filepath, "w", encoding="utf-8") as f:
                f.write(content)
            return f"✅ File created: {filepath} (as text)"
    
    except Exception as e:
        return f"❌ Error creating document: {e}"


# ============================================
# TOOL 5: APP OPENER
# ============================================

APP_ALIASES = {
    # Browsers — use 'start' with protocol or known paths
    "browser": 'start "" "https://www.google.com"',
    "chrome": 'start "" "chrome"',
    "google chrome": 'start "" "chrome"',
    "firefox": 'start "" "firefox"',
    "edge": 'start "" "msedge"',
    "brave": 'start "" "brave"',
    # Code editors
    "vscode": "code",
    "vs code": "code",
    "visual studio code": "code",
    "notepad": "notepad",
    "notepad++": 'start "" "notepad++"',
    # System tools
    "calculator": "calc",
    "calc": "calc",
    "terminal": "start cmd",
    "cmd": "start cmd",
    "powershell": "start powershell",
    "explorer": "explorer",
    "file explorer": "explorer",
    "task manager": "taskmgr",
    "settings": "start ms-settings:",
    "control panel": "control",
    # Apps — use 'start' for Windows Store / installed apps
    "spotify": 'start "" "spotify"',
    "discord": 'start "" "discord"',
    "slack": 'start "" "slack"',
    "whatsapp": 'start "" "whatsapp"',
    "telegram": 'start "" "telegram"',
    "obs": 'start "" "obs64"',
    "paint": "mspaint",
    "word": 'start "" "winword"',
    "excel": 'start "" "excel"',
    "powerpoint": 'start "" "powerpnt"',
}

# Common browser URL patterns
WEB_KEYWORDS = ["http://", "https://", "www.", ".com", ".org", ".net", ".io", ".dev"]


def app_opener(args):
    """
    Open an application by name or a URL in the default browser.
    Args: {"app": "vscode", "path": "optional/file/to/open", "url": "optional_url"}
    """
    app_name = args.get("app", "").strip().lower()
    open_path = args.get("path", "").strip()
    url = args.get("url", "").strip()

    # If URL provided, open in default browser
    if url:
        if not url.startswith(("http://", "https://")):
            url = "https://" + url
        try:
            subprocess.Popen(f'start "" "{url}"', shell=True)
            return f"✅ Opened URL: {url}"
        except Exception as e:
            return f"❌ Failed to open URL: {e}"

    if not app_name:
        return "❌ No application name provided."

    # Check if app_name is actually a URL
    if any(kw in app_name for kw in WEB_KEYWORDS):
        url = app_name if app_name.startswith(("http://", "https://")) else "https://" + app_name
        try:
            subprocess.Popen(f'start "" "{url}"', shell=True)
            return f"✅ Opened URL: {url}"
        except Exception as e:
            return f"❌ Failed to open URL: {e}"

    # Resolve alias
    cmd = APP_ALIASES.get(app_name, None)

    if cmd is None:
        # Try using 'start' as a generic launcher (Windows will find it)
        cmd = f'start "" "{app_name}"'

    # Add path if provided
    if open_path:
        if "code" in cmd:
            cmd = f'code "{open_path}"'
        else:
            cmd = f'{cmd} "{open_path}"'

    try:
        subprocess.Popen(cmd, shell=True)
        return f"✅ Opened: {app_name}" + (f" with {open_path}" if open_path else "")
    except Exception as e:
        return f"❌ Failed to open {app_name}: {e}"


# ============================================
# TOOL 6: SYSTEM INFO
# ============================================

def system_info(args):
    """
    Get system information.
    Args: {"query": "os|disk|processes|cpu|ram|battery|ip|env|all"}
    """
    query = args.get("query", "all").lower()
    info_parts = []

    if query in ("os", "all"):
        info_parts.append(
            f"🖥️ OS: {platform.system()} {platform.release()} ({platform.machine()})\n"
            f"   Computer: {platform.node()}\n"
            f"   Python: {sys.version.split()[0]}"
        )

    if query in ("cpu", "all"):
        info_parts.append(
            f"⚡ CPU: {psutil.cpu_percent()}% usage ({psutil.cpu_count()} cores)"
        )

    if query in ("ram", "memory", "all"):
        ram = psutil.virtual_memory()
        info_parts.append(
            f"🧠 RAM: {ram.percent}% used "
            f"({ram.used / (1024**3):.1f} GB / {ram.total / (1024**3):.1f} GB)"
        )

    if query in ("disk", "all"):
        for part in psutil.disk_partitions():
            try:
                usage = psutil.disk_usage(part.mountpoint)
                info_parts.append(
                    f"💾 {part.device}: {usage.percent}% used "
                    f"({usage.used / (1024**3):.1f} GB / {usage.total / (1024**3):.1f} GB)"
                )
            except PermissionError:
                pass

    if query in ("battery", "all"):
        battery = psutil.sensors_battery()
        if battery:
            info_parts.append(
                f"🔋 Battery: {battery.percent}% "
                f"({'Charging' if battery.power_plugged else 'Discharging'})"
            )

    if query in ("processes", "procs"):
        procs = []
        for p in psutil.process_iter(['pid', 'name', 'cpu_percent']):
            try:
                procs.append(p.info)
            except (psutil.NoSuchProcess, psutil.AccessDenied):
                pass
        # Top 10 by CPU
        procs.sort(key=lambda x: x.get('cpu_percent', 0) or 0, reverse=True)
        lines = ["📊 Top processes:"]
        for p in procs[:10]:
            lines.append(f"  {p['pid']:>6} | {p['name']:<30} | CPU: {p.get('cpu_percent', 0):.1f}%")
        info_parts.append("\n".join(lines))

    if query in ("ip", "network"):
        try:
            import socket
            hostname = socket.gethostname()
            local_ip = socket.gethostbyname(hostname)
            info_parts.append(f"🌐 Host: {hostname}\n   Local IP: {local_ip}")
        except Exception:
            info_parts.append("🌐 Could not determine IP")

    if query in ("env",):
        important_vars = ["PATH", "USERPROFILE", "HOME", "COMPUTERNAME", "OS"]
        lines = ["🔧 Environment:"]
        for var in important_vars:
            val = os.environ.get(var, "not set")
            if len(val) > 100:
                val = val[:100] + "..."
            lines.append(f"  {var}: {val}")
        info_parts.append("\n".join(lines))

    if not info_parts:
        return f"❌ Unknown query: {query}. Use: os, disk, cpu, ram, battery, processes, ip, env, all"

    return "\n\n".join(info_parts)


# ============================================
# TOOL 7: MEMORY STORE
# ============================================

def memory_store(args):
    """
    Persistent memory operations.
    Args: {"action": "save|get|search|list|delete|clear",
           "key": "...", "value": "..."}
    """
    action = args.get("action", "").lower()

    if action == "save":
        key = args.get("key", "")
        value = args.get("value", "")
        if not key:
            return "❌ No key provided."
        return mem.save(key, value)

    elif action == "get":
        key = args.get("key", "")
        if not key:
            return "❌ No key provided."
        val = mem.get(key)
        if val is None:
            return f"❌ Memory '{key}' not found."
        return f"🧠 {key}: {val}"

    elif action == "search":
        query = args.get("query", args.get("key", ""))
        if not query:
            return "❌ No search query provided."
        return mem.search(query)

    elif action == "list":
        return mem.list_all()

    elif action == "delete":
        key = args.get("key", "")
        if not key:
            return "❌ No key provided."
        return mem.delete(key)

    elif action == "clear":
        return mem.clear_all()

    else:
        return f"❌ Unknown action: {action}. Use: save, get, search, list, delete, clear"


# ============================================
# TOOL REGISTRY
# ============================================

TOOL_REGISTRY = {
    "terminal_executor": {
        "func": terminal_executor,
        "description": "ALWAYS use this to run any shell/terminal command. Args: {command, timeout?, cwd?}",
    },
    "file_manager": {
        "func": file_manager,
        "description": "ALWAYS use this for ANY file operation: create, read, write, delete, move, copy, list, search, mkdir. Relative paths resolve to WORKSPACE. Args: {action, path, content?, destination?, query?}",
    },
    "web_search": {
        "func": web_search,
        "description": "ALWAYS use this to search the web for real-time information using DuckDuckGo. Args: {query}",
    },
    "document_creator": {
        "func": document_creator,
        "description": "ALWAYS use this to create documents (txt, md, html, py, docx). Supports markdown formatting in docx. Default save to WORKSPACE. Args: {filename, content, format?}",
    },
    "app_opener": {
        "func": app_opener,
        "description": "ALWAYS use this to open ANY application, website, or software (chrome, whatsapp, vscode, spotify, ANY app). Args: {app, path?, url?}",
    },
    "system_info": {
        "func": system_info,
        "description": "ALWAYS use this for system information: os, disk, cpu, ram, battery, processes, ip, env, all. Args: {query}",
    },
    "memory_store": {
        "func": memory_store,
        "description": "ALWAYS use this to save/retrieve user memories and preferences. Args: {action, key?, value?, query?}",
    },
}


def get_tool_descriptions():
    """Return formatted tool list for the system prompt."""
    lines = []
    for i, (name, tool) in enumerate(TOOL_REGISTRY.items(), 1):
        lines.append(f"{i}. {name}: {tool['description']}")
    return "\n".join(lines)


def execute_tool(name, args):
    """Execute a tool by name with arguments dict."""
    if name not in TOOL_REGISTRY:
        return f"❌ Tool '{name}' not found. Available: {', '.join(TOOL_REGISTRY.keys())}"

    func = TOOL_REGISTRY[name]["func"]

    try:
        return func(args)
    except Exception as e:
        return f"❌ Tool '{name}' error: {e}"
